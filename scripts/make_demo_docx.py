#!/usr/bin/env python3
"""Generates the demo and submission document, docs/Claims-Assistant-Demo.docx.

    uv run --with python-docx python scripts/make_demo_docx.py
    uv run --with python-docx python scripts/make_demo_docx.py --author "Your Name" --out ~/Desktop/demo.docx

Without uv: pip install python-docx, then python3 scripts/make_demo_docx.py.
Options: --author, --live-url, --repo-url, --image, --out (see --help)."""

import argparse
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description="Generate the Claims Assistant demo document (.docx).")
parser.add_argument("--author", default="HarshithKoriRaj", help="name shown as the submitter")
parser.add_argument("--live-url", default="https://claims-assistant-4wy7.onrender.com", help="deployed demo URL")
parser.add_argument("--repo-url", default="https://github.com/HarshithKoriRaj/insurance-claims-sop-harness", help="source repository URL")
parser.add_argument("--image", default="ghcr.io/harshithkoriraj/insurance-claims-sop-harness:latest", help="published Docker image")
parser.add_argument("--out", type=Path, default=ROOT / "docs" / "Claims-Assistant-Demo.docx", help="output .docx path")
args = parser.parse_args()
OUT = args.out.expanduser()
LIVE = args.live_url.rstrip("/")
REPO = args.repo_url.rstrip("/")
IMAGE = args.image
INDIGO = RGBColor(0x31, 0x2E, 0x81)

doc = Document()
styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)
for name in ("Heading 1", "Heading 2"):
    styles[name].font.color.rgb = INDIGO


def hyperlink(paragraph, url, text=None):
    rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1D4ED8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text or url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def bullet(text, bold_lead=None, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
        p.add_run(" " + text)
    else:
        p.add_run(text)
    return p


def table(rows, header=True, widths=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(9.5)
            if header and r == 0:
                run.bold = True
    doc.add_paragraph()
    return t


def code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    return p


# ------------------------------------------------------------------ title
title = doc.add_heading("Claims Assistant", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
sub = doc.add_paragraph()
sub.add_run("An SOP harness for an insurance claims support chat agent — demo and submission notes").italic = True
doc.add_paragraph(f"Submitted by {args.author} · {date.today():%B %Y}")

# ------------------------------------------------------------------ links
doc.add_heading("Submission links", level=1)
links = doc.add_table(rows=4, cols=2)
links.style = "Light Grid Accent 1"
entries = [
    ("Live demo", LIVE),
    ("Source code and setup docs", REPO),
    ("Docker image", None),
    ("End-to-end demo story", REPO + "/blob/main/docs/demo-story.md"),
]
for row, (label, url) in enumerate(entries):
    links.cell(row, 0).text = ""
    links.cell(row, 0).paragraphs[0].add_run(label).bold = True
    cell = links.cell(row, 1)
    cell.text = ""
    if url:
        hyperlink(cell.paragraphs[0], url)
    else:
        run = cell.paragraphs[0].add_run(f"docker run -p 8000:8000 -e OPENAI_API_KEY=… {IMAGE}")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
doc.add_paragraph()
note = doc.add_paragraph()
note.add_run("Note: ").bold = True
note.add_run("the live demo runs on Render's free plan. After 15 minutes without visitors it sleeps, and the first page load "
             "takes about a minute to wake it. Everything after that is normal speed.")

# ------------------------------------------------------------------ what it does
doc.add_heading("What it does", level=1)
doc.add_paragraph(
    "The assistant follows a four-phase standard operating procedure. A language model reads each message and words each "
    "reply; deterministic code decides everything in between: who the caller is, what they may see, and what happens next."
)
table([
    ["Phase", "What happens", "Gate to leave it"],
    ["VERIFY_ID", "Collects identity details; discloses nothing about claims; handles partial answers, refusals and clarifying questions",
     "3 of 5 permitted fields (name, date of birth, phone, email, SSN last 4) match exactly one policyholder, with no contradiction"],
    ["RESOLVE_INTENT", "Works out which claim the caller means, using hints remembered from before verification",
     "Exactly one of the caller's own claims selected"],
    ["PROCESS_CASE", "Answers only from the claim record and approved guidance", "The caller is done, or picks another claim"],
    ["POST_PROCESS", "Shows an email summary and asks the caller to send or skip it", "An explicit send or skip; nothing is emailed without send"],
])

# ------------------------------------------------------------------ try it
doc.add_heading("Try it in two minutes", level=1)
for step in [
    "Open the live demo and click the “Demo: Margaret Chen” suggestion above the message box, then Send.",
    "Expect: verified from name, date of birth and SSN last 4 (the policy number doesn't count); routed straight to "
    "claim CL-2048 from her remembered hint, without being asked which claim; the denial reason, the documents needed, "
    "and the appeal deadline, 17 days from the pinned demo date of March 1, 2026.",
    "Click “I'm done”, then Send. The summary card shows the email preview; after Send summary it is recorded "
    "(the live demo has no mail server; with Docker Compose it arrives in the Mailpit inbox).",
    "Try “What is RL?” three times, a wrong SSN three times, or “Talk to a human” to see the recovery rules.",
]:
    bullet(step, style="List Number")
doc.add_paragraph("To check everything automatically — five conversations, 32 checks, exit code 0 when all pass:")
code(f"python3 scripts/story.py {LIVE}")

# ------------------------------------------------------------------ architecture
doc.add_heading("Architecture", level=1)
doc.add_paragraph("Each chat turn runs the same fixed pipeline:")
for step in [
    ("Interpret.", "One model call with a forced function call returns typed labels: identity details as stated, intent, "
     "claim hints, scope, emotion, and whether the caller wants a human, refuses, or chooses send/skip."),
    ("Ground.", "Code keeps an identity value or claim hint only if it actually appears in the caller's message, so the "
     "model can't invent a year or complete an email address."),
    ("Decide.", "The deterministic workflow engine applies the labels, runs verification, looks up claims scoped to the "
     "verified caller, computes deadlines, picks approved guidance, and makes every phase transition."),
    ("Respond.", "One model call writes the reply from the engine's brief, which contains only facts the caller is allowed "
     "to see. Before verification the brief holds no claim data at all, and a backstop discards any reply that mentions "
     "a claim ID or an amount."),
]:
    bullet(step[1], bold_lead=step[0], style="List Number")
doc.add_paragraph(
    "Stack: React + TypeScript UI, FastAPI (Python 3.13), Pydantic contracts, SQLite sessions, OpenAI (gpt-4.1-mini by "
    "default) or Claude, Docker, GitHub Actions. If the model is unavailable, interpretation falls back to deterministic "
    "rules and replies to templates, so the gates keep working."
)

# ------------------------------------------------------------------ faster
doc.add_heading("Why this architecture is faster", level=1)
doc.add_paragraph(
    "The model is used only for language. Every decision — identity, ownership, dates, consent, transitions — is ordinary "
    "code, so each turn costs a fixed, small number of short model calls instead of an open-ended agent loop."
)
for lead, text in [
    ("A fixed two-call turn.",
     "Every message costs exactly two model calls (interpret, respond), however many checks and lookups run. A typical "
     "tool-calling agent needs a model round trip per tool: verify the caller, list claims, read the claim, read guidance, "
     "then answer — about five sequential calls, each re-sending a growing context."),
    ("Decisions in microseconds, not model round trips.",
     "Identity normalization and matching, claim lookup, deadline arithmetic, guidance selection and the consent gate run as "
     "in-memory code. Fixtures are validated and indexed once at startup."),
    ("Fewer calls where none are needed.",
     "Starting a conversation uses a template (no model call); the Send, Skip, End and Human buttons skip interpretation "
     "(one call instead of two)."),
    ("Short prompts and a small model.",
     "The interpret call sees only the latest message and a little context; the reply call sees the brief and the last "
     "12 turns. Because policy lives in code, a small fast model is enough."),
    ("No repair loops.",
     "Structured output comes from a forced function call; invalid or ungrounded output is dropped or replaced by the rule "
     "reading, never re-prompted."),
    ("Faster to change safely.",
     "301 deterministic tests run in about 2 seconds without any model; thresholds live in a policy file; CI checks every push."),
]:
    bullet(text, bold_lead=lead)
doc.add_paragraph("Measured with gpt-4.1-mini, on the live demo (Render free plan) and locally:")
table([
    ["Interaction", "Model calls", "Observed"],
    ["Start a conversation", "0", "About 25 ms"],
    ["Button action (send, skip, end, human)", "1", "About 1 s"],
    ["Chat message", "2", "About 1.5–4.5 s"],
    ["Full test suite (301 tests, no model)", "0", "About 2 s"],
])

# ------------------------------------------------------------------ Jev
doc.add_heading("Adapting Jev for faster classification", level=1)
doc.add_paragraph(
    "Jev is TypeSafe AI's “System One” classification model, released 15 September 2026 (early access). It answers typed "
    "questions — Choice, Score, and Noul (probability a statement is true) — with calibrated probabilities, and never generates "
    "text. Questions in one request run in parallel. TypeSafe reports 70–500 ms end to end at $0.042 per million input tokens "
    "(output is free), which is well under a generative model call."
)
doc.add_heading("Where it plugs in", level=2)
doc.add_paragraph(
    "The architecture already separates labeling from deciding, so Jev slots into the Interpret step without touching the "
    "engine or its gates. One Jev request per turn asks the current phase's questions in parallel, alongside the extraction "
    "call; a label is used only at or above its threshold, otherwise the model's label (today's path) applies."
)
table([
    ["Question", "Type", "Act at or above", "Otherwise"],
    ["scope (claim support, general, privacy, small talk, unrelated…)", "Choice", "0.80", "Treat as in scope and ask what they need"],
    ["emotion", "Choice", "0.70", "Neutral, respectful wording"],
    ["wants_human", "Noul", "0.85", "0.40–0.85: ask whether they want a person"],
    ["refuses_step", "Noul", "0.80", "Explain why the step matters, ask once more"],
    ["intent", "Choice", "0.70", "Ask one clarifying question"],
    ["followup_topic", "Choice", "0.70", "Ask which topic they mean"],
    ["summary_decision", "Choice", "send 0.95 / skip 0.80", "Ask yes or no; never send"],
    ["done", "Noul", "0.85", "Ask whether there is anything else"],
])
doc.add_paragraph("Thresholds are starting values, to be replaced by what an evaluation set supports.")
doc.add_heading("What stays out of Jev, by design", level=2)
for lead, text in [
    ("Values and arithmetic.", "Identity values, dates, amounts and counts stay with the extraction call and deterministic code — "
     "TypeSafe documents that Jev 1.13 reads dates as text and doesn't count or do arithmetic."),
    ("Every gate.", "Jev never decides identity, ownership, permissions or phase transitions. A send label is valid only while "
     "exactly one summary offer is active; below 0.95 the assistant asks yes or no."),
    ("Privacy.", "Only the latest message is sent, after deterministic redaction (emails, phone numbers, digit runs, dates become "
     "[EMAIL], [PHONE], [DIGITS], [DATE]), with the phase name. No claim facts, stored personal data or older turns."),
    ("Resilience.", "Jev is optional. On a timeout, an error or no key, the extraction call's labels are used and the turn records "
     "which classifier answered."),
]:
    bullet(text, bold_lead=lead)
doc.add_heading("Expected gain and status", level=2)
doc.add_paragraph(
    "With labels from Jev in 70–500 ms, the interpret step can shrink to a small extraction-only call — or be skipped entirely "
    "for messages with no identity details. We estimate that would cut roughly a third to half of each turn's latency and "
    "most of its classification cost (an estimate from TypeSafe's reported latency, not yet measured here). Confidence "
    "thresholds also make each routing decision measurable and auditable."
)
status = doc.add_paragraph()
status.add_run("Status: ").bold = True
status.add_run(
    "the Jev integration is fully designed (docs/2026-09-28-sop-harness-plan.md, “Fast classification with Jev”) but was "
    "scoped out of the submitted build to meet the deadline. The code is ready for it: a Jev classifier would fill the label "
    "fields of the existing Interpretation contract behind the same provider interface, configured with TYPESAFE_API_KEY and "
    "the pinned model jev-1.13.0."
)

# ------------------------------------------------------------------ safety and verification
doc.add_heading("Safety and verification", level=1)
for lead, text in [
    ("No disclosure before verification;", "claim lookups require a verified party ID, and another customer's claim looks exactly "
     "like a missing one."),
    ("Consent before email;", "the full preview is shown, it goes only to the address on file, and each preview sends at most once "
     "(even under a double click)."),
    ("Nothing guessed:", "an ambiguous date or a mistyped phone number is asked about, and doesn't count as a failed attempt; three "
     "failed attempts stop automated verification."),
    ("Public-demo guards:", "per-client rate limits protect the model budget; keys live only in environment secrets, and the full git "
     "history was scanned before the repository went public."),
    ("Tests:", "301 automated tests (offline, deterministic), 6 live-model scenario tests, and the 32-check end-to-end story; "
     "CI runs on every push and publishes the Docker image."),
]:
    bullet(text, bold_lead=lead)

doc.add_heading("Known limitations", level=1)
for text in [
    "Human handoff is simulated: the conversation moves to a handoff state and never claims a person has joined.",
    "The live demo has no mail server, so approved summaries are recorded, not delivered; Docker Compose includes the Mailpit inbox.",
    "Someone calling on a policyholder's behalf is offered a human instead of an automated authorization flow.",
    "Sessions use SQLite and reset on the free plan's restarts; the claim data is the supplied static fixtures.",
]:
    bullet(text)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"wrote {OUT}")
